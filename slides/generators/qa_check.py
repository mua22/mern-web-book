"""Reusable QA for generated lecture decks: geometry (off-slide/overlap) +
placeholder-text scan, using python-pptx (no LibreOffice/markitdown needed).

Usage: python qa_check.py <deck.pptx> [--dump]
"""
import sys
import re
from pptx import Presentation

SW, SH = 13.333, 7.5
MARGIN = 0.35


def emu_in(v):
    return v / 914400.0 if v is not None else None


def main():
    path = sys.argv[1]
    dump = "--dump" in sys.argv
    prs = Presentation(path)
    issues = []
    placeholder_re = re.compile(r"\bxxx\b|lorem ipsum|\bTODO\b|\[insert|undefined|NaN", re.I)

    for i, slide in enumerate(prs.slides, 1):
        boxes = []
        for sh in slide.shapes:
            try:
                x, y, w, h = emu_in(sh.left), emu_in(sh.top), emu_in(sh.width), emu_in(sh.height)
            except Exception:
                x = y = w = h = None
            txt = ""
            if sh.has_text_frame:
                txt = " / ".join(p.text for p in sh.text_frame.paragraphs if p.text)
            if dump and txt:
                print(f"  [S{i}] ({x:.2f},{y:.2f} {w:.2f}x{h:.2f})  {txt[:100]!r}")
            if txt and placeholder_re.search(txt):
                issues.append(f"S{i}: placeholder-looking text: {txt[:80]!r}")
            if None not in (x, y, w, h):
                if x < -0.05 or y < -0.05 or x + w > SW + 0.05 or y + h > SH + 0.05:
                    issues.append(f"S{i}: OFFSLIDE ({x:.2f},{y:.2f})-({x+w:.2f},{y+h:.2f}) {txt[:50]!r}")
                if txt:
                    boxes.append((x, y, w, h, txt))
        for a in range(len(boxes)):
            for b in range(a + 1, len(boxes)):
                ax, ay, aw, ah, at = boxes[a]
                bx, by, bw, bh, bt = boxes[b]
                ox = max(0, min(ax + aw, bx + bw) - max(ax, bx))
                oy = max(0, min(ay + ah, by + bh) - max(ay, by))
                if ox > 0.15 and oy > 0.15 and (ox * oy) > 0.3:
                    issues.append(f"S{i}: OVERLAP {at[:30]!r} <> {bt[:30]!r} ({ox:.2f}x{oy:.2f})")

    print(f"\n{len(prs.slides.__iter__().__length_hint__() if hasattr(prs.slides, '__length_hint__') else list(prs.slides))} -- checked {path}")
    if issues:
        print(f"{len(issues)} ISSUE(S):")
        for iss in issues:
            print("  " + iss)
    else:
        print("No geometry/placeholder issues found.")


if __name__ == "__main__":
    main()
